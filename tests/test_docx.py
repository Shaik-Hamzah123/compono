"""Golden/invariant tests for compono.docx: render a spec, reopen the .docx
with python-docx, and assert real structural properties.

Covers the docx-specific invariant: every primitive except `chart` (which is
documented as a rasterized image — no native Word chart API exists) produces
a genuine, editable python-docx object.
"""

import json
from pathlib import Path

import docx as python_docx
import pytest

from compono.cli import main as cli_main
from compono.docx import (
    DocxRenderReport,
    DocxValidationError,
    render_docx,
    validate_docx,
)

FULL_SPEC = {
    "title": "Training Proposal",
    "sections": [
        {
            "header_text": "Confidential",
            "footer_text": "Page footer",
            "body": [
                {"primitive": "heading", "text": "Overview", "level": 1},
                {
                    "primitive": "paragraph",
                    "runs": [
                        {"text": "This is "},
                        {"text": "bold", "bold": True},
                        {"text": " and this is a "},
                        {"text": "link", "link": "https://example.com"},
                        {"text": "."},
                    ],
                },
                {
                    "primitive": "bullet_list",
                    "items": [[{"text": "First point"}], [{"text": "Second point"}]],
                },
                {
                    "primitive": "numbered_list",
                    "items": [[{"text": "Step one"}], [{"text": "Step two"}]],
                },
                {
                    "primitive": "table",
                    "headers": ["Track", "Weeks"],
                    "rows": [["AI Foundations", "1-2"], ["Advanced ML", "3-4"]],
                },
                {"primitive": "image", "placeholder": True, "caption": "Company logo"},
                {
                    "primitive": "chart",
                    "chart_type": "bar",
                    "categories": ["Q1", "Q2", "Q3"],
                    "series": [{"name": "Revenue", "values": [10, 20, 15]}],
                },
                {"primitive": "page_break"},
            ],
        }
    ],
}


def test_validate_docx_accepts_valid_spec() -> None:
    report = validate_docx(FULL_SPEC)
    assert report.valid is True
    assert report.errors == []


def test_validate_docx_returns_structured_errors_for_malformed_spec() -> None:
    report = validate_docx(
        {"title": "Bad", "sections": [{"body": [{"primitive": "heading"}]}]}
    )
    assert report.valid is False
    error = report.errors[0]
    assert set(error.keys()) >= {
        "section",
        "primitive",
        "field",
        "error",
        "detail",
        "fix",
    }


def test_render_docx_writes_a_real_docx_file(tmp_path: Path) -> None:
    output = tmp_path / "report.docx"
    report = render_docx(FULL_SPEC, output)

    assert isinstance(report, DocxRenderReport)
    assert report.docx_path == output
    assert output.exists()


def test_render_docx_raises_structured_error_and_writes_nothing(tmp_path: Path) -> None:
    output = tmp_path / "report.docx"
    with pytest.raises(DocxValidationError) as exc_info:
        render_docx(
            {"title": "Bad", "sections": [{"body": [{"primitive": "heading"}]}]}, output
        )

    assert exc_info.value.errors
    assert not output.exists()


def test_heading_and_paragraph_text_round_trip_as_real_text(tmp_path: Path) -> None:
    output = tmp_path / "report.docx"
    render_docx(FULL_SPEC, output)

    doc = python_docx.Document(str(output))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Overview" in all_text
    assert "This is bold and this is a link." in all_text


def test_bold_and_link_runs_are_styled(tmp_path: Path) -> None:
    output = tmp_path / "report.docx"
    render_docx(FULL_SPEC, output)

    doc = python_docx.Document(str(output))
    body_para = next(p for p in doc.paragraphs if "This is" in p.text)
    bold_runs = [r for r in body_para.runs if r.bold]
    assert bold_runs and bold_runs[0].text == "bold"

    # The hyperlink is built via raw XML (w:hyperlink), not a plain w:r run —
    # confirm it exists on the paragraph's underlying XML.
    hyperlink_els = body_para._p.findall(
        ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink"
    )
    assert len(hyperlink_els) == 1


def test_bullet_and_numbered_lists_use_real_list_styles(tmp_path: Path) -> None:
    output = tmp_path / "report.docx"
    render_docx(FULL_SPEC, output)

    doc = python_docx.Document(str(output))
    bullet_paragraphs = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    numbered_paragraphs = [p for p in doc.paragraphs if p.style.name == "List Number"]
    assert [p.text for p in bullet_paragraphs] == ["First point", "Second point"]
    assert [p.text for p in numbered_paragraphs] == ["Step one", "Step two"]


def test_table_data_round_trips_into_a_real_table(tmp_path: Path) -> None:
    output = tmp_path / "report.docx"
    render_docx(FULL_SPEC, output)

    doc = python_docx.Document(str(output))
    table = doc.tables[0]
    header_row = [table.cell(0, c).text for c in range(2)]
    assert header_row == ["Track", "Weeks"]
    assert table.cell(1, 0).text == "AI Foundations"
    assert table.cell(2, 1).text == "3-4"


def test_image_placeholder_renders_as_text_not_a_picture(tmp_path: Path) -> None:
    output = tmp_path / "report.docx"
    render_docx(FULL_SPEC, output)

    doc = python_docx.Document(str(output))
    # FULL_SPEC also has a chart primitive, which legitimately renders as one
    # real inline picture — the placeholder itself must not add another.
    assert len(doc.inline_shapes) == 1
    placeholder_text = " ".join(p.text for p in doc.paragraphs)
    assert "Company logo" in placeholder_text


def test_image_placeholder_appears_in_manifest(tmp_path: Path) -> None:
    output = tmp_path / "report.docx"
    report = render_docx(FULL_SPEC, output)

    assert any(entry.get("caption") == "Company logo" for entry in report.manifest)


def test_chart_renders_as_an_embedded_picture(tmp_path: Path) -> None:
    output = tmp_path / "report.docx"
    render_docx(FULL_SPEC, output)

    doc = python_docx.Document(str(output))
    # The chart is the only real picture in this spec (the image primitive
    # above is a placeholder, not a picture) — documented rasterization.
    assert len(doc.inline_shapes) == 1


def test_section_header_and_footer_text_are_applied(tmp_path: Path) -> None:
    output = tmp_path / "report.docx"
    render_docx(FULL_SPEC, output)

    doc = python_docx.Document(str(output))
    section = doc.sections[0]
    assert section.header.paragraphs[0].text == "Confidential"
    assert section.footer.paragraphs[0].text == "Page footer"


def test_multiple_sections_produce_multiple_docx_sections(tmp_path: Path) -> None:
    spec = {
        "title": "Two Sections",
        "sections": [
            {
                "header_text": "Section 1",
                "body": [{"primitive": "heading", "text": "A", "level": 1}],
            },
            {
                "header_text": "Section 2",
                "body": [{"primitive": "heading", "text": "B", "level": 1}],
            },
        ],
    }
    output = tmp_path / "two.docx"
    render_docx(spec, output)

    doc = python_docx.Document(str(output))
    assert len(doc.sections) == 2
    assert doc.sections[0].header.paragraphs[0].text == "Section 1"
    assert doc.sections[1].header.paragraphs[0].text == "Section 2"


def test_cli_docx_validate_and_render_smoke(tmp_path: Path) -> None:
    spec_path = tmp_path / "spec.json"
    spec_path.write_text(json.dumps(FULL_SPEC), encoding="utf-8")

    assert cli_main(["docx", "validate", str(spec_path)]) == 0

    out_path = tmp_path / "out.docx"
    exit_code = cli_main(["docx", "render", str(spec_path), "-o", str(out_path)])
    assert exit_code == 0
    assert out_path.exists()
