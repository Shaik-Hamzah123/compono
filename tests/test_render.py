"""Unit/integration tests for src/compono/render.py and cli.py.

Covers the two public verbs (render_deck/validate) end-to-end for the v1
primitive slice (header, text, grid, shape), plus a CLI smoke test against
an inline minimal spec.
"""

import dataclasses
import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE

from compono.cli import main as cli_main
from compono.render import (
    DeckValidationError,
    _extract_text_fields,
    _resolve_font_path,
    _sequence_step_rects,
    _table_cell_rects,
    render_deck,
    validate,
)
from compono.resolver import Rect, Template
from compono.schema import Sequence, SequenceStep, Table

EXAMPLES_DIR = Path(__file__).parent.parent / "examples"
FULL_CATALOG_SPEC = json.loads(
    (EXAMPLES_DIR / "full_catalog.json").read_text(encoding="utf-8")
)

# Kept inline (not examples/minimal.json, which was removed as a showcase
# example — every example deck is now real reference material, not a bare
# smoke-test fixture) so this suite has no dependency on the examples/ dir
# beyond the genuinely representative decks.
MINIMAL_SPEC = {
    "slides": [
        {
            "header": {
                "title": "Welcome to compono",
                "subtitle": "Agent-oriented, code-based PPTX generation",
                "eyebrow": "Demo",
            },
            "body": [
                {
                    "primitive": "text",
                    "mode": "bullets",
                    "content": [
                        "The agent never writes raw x/y/w/h coordinates",
                        "A constraint-based resolver computes real EMU positions",
                        "Every primitive renders as a genuine, editable OOXML shape",
                    ],
                    "emphasis_indices": [2],
                },
                {
                    "primitive": "grid",
                    "columns": 2,
                    "items": [
                        {
                            "primitive": "shape",
                            "id": "box-a",
                            "kind": "rounded_rect",
                            "fill": "#4F46E5",
                            "text": {"content": "Box A"},
                        },
                        {
                            "primitive": "shape",
                            "id": "box-b",
                            "kind": "rounded_rect",
                            "fill": "#059669",
                            "text": {"content": "Box B"},
                        },
                    ],
                },
                {
                    "primitive": "shape",
                    "kind": "connector",
                    "connects": {"from_id": "box-a", "to_id": "box-b"},
                },
            ],
        }
    ]
}


def test_validate_accepts_minimal_spec() -> None:
    report = validate(MINIMAL_SPEC)
    assert report.valid is True
    assert report.errors == []


def test_validate_rejects_malformed_spec() -> None:
    report = validate(
        {"slides": [{"body": [{"primitive": "header"}]}]}
    )  # missing required title
    assert report.valid is False
    assert report.errors[0]["error"] is not None


def test_validate_flags_unknown_connector_id() -> None:
    spec = {
        "slides": [
            {
                "body": [
                    {
                        "primitive": "shape",
                        "kind": "connector",
                        "connects": {"from_id": "x", "to_id": "y"},
                    },
                ]
            }
        ]
    }
    report = validate(spec)
    assert report.valid is False
    assert any(e["error"] == "layout" for e in report.errors)


def test_render_deck_writes_a_real_pptx(tmp_path: Path) -> None:
    output = tmp_path / "deck.pptx"
    report = render_deck(MINIMAL_SPEC, output)

    assert report.pptx_path == output
    assert output.exists()
    assert len(report.actual_layout) == 1

    prs = Presentation(str(output))
    slides = list(prs.slides)
    assert len(slides) == 1

    slide = slides[0]
    # Non-negotiable invariant: every primitive is a real, editable shape —
    # never a picture or embedded video (COMPONO_PLAN.md section 3).
    shape_types = {shape.shape_type for shape in slide.shapes}

    assert MSO_SHAPE_TYPE.PICTURE not in shape_types
    # header textbox + bullets textbox + 2 shapes + 1 connector + footer = 6 shapes
    assert len(slide.shapes) == 6


def test_render_deck_raises_on_invalid_spec(tmp_path: Path) -> None:
    with pytest.raises(DeckValidationError):
        render_deck(
            {"slides": [{"body": [{"primitive": "header"}]}]}, tmp_path / "out.pptx"
        )


def test_render_deck_respects_explicit_template(tmp_path: Path) -> None:
    template = Template.from_yaml()
    output = tmp_path / "deck.pptx"
    render_deck(MINIMAL_SPEC, output, template=template)
    prs = Presentation(str(output))
    assert prs.slide_width == template.page_width
    assert prs.slide_height == template.page_height


def test_render_deck_uses_the_font_family_named_by_deck_template(
    tmp_path: Path,
) -> None:
    """`Deck.template` (schema.py) was previously unwired — every render used
    the hardcoded default regardless of what the spec asked for.
    """
    spec = {**MINIMAL_SPEC, "template": "modern"}
    output = tmp_path / "deck.pptx"
    render_deck(spec, output)

    prs = Presentation(str(output))
    fonts = {
        p.font.name
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
        for p in shape.text_frame.paragraphs
    }
    assert fonts == {"Georgia"}


def test_render_deck_raises_on_unknown_template_name(tmp_path: Path) -> None:
    spec = {**MINIMAL_SPEC, "template": "nonexistent"}
    with pytest.raises(DeckValidationError) as exc_info:
        render_deck(spec, tmp_path / "out.pptx")
    assert exc_info.value.errors[0]["error"] == "unknown_template"


def test_validate_reports_unknown_template_name_as_a_structured_error() -> None:
    spec = {**MINIMAL_SPEC, "template": "nonexistent"}
    report = validate(spec)
    assert report.valid is False
    assert report.errors[0]["error"] == "unknown_template"


def test_explicit_template_kwarg_overrides_deck_template_field(
    tmp_path: Path,
) -> None:
    spec = {**MINIMAL_SPEC, "template": "modern"}
    output = tmp_path / "deck.pptx"
    render_deck(spec, output, template=Template.from_yaml())  # "default"

    prs = Presentation(str(output))
    fonts = {
        p.font.name
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
        for p in shape.text_frame.paragraphs
    }
    assert fonts == {"Calibri"}


def test_render_shape_text_applies_color_and_font_to_every_line(
    tmp_path: Path,
) -> None:
    """shape.text.content with embedded "\\n"s becomes multiple paragraphs
    (python-pptx's text_frame.text setter splits on newline) — a bug once
    only formatted paragraphs[0], leaving every line after the first in the
    theme default font/color instead of the one the agent asked for.
    """
    spec = {
        "slides": [
            {
                "header": {"title": "Card"},
                "body": [
                    {
                        "primitive": "shape",
                        "kind": "rounded_rect",
                        "fill": "#FDB71A",
                        "text": {
                            "content": "Line one\nLine two\nLine three",
                            "color": "#1A2332",
                        },
                    }
                ],
            }
        ]
    }
    output = tmp_path / "deck.pptx"
    render_deck(spec, output)

    prs = Presentation(str(output))
    shape = next(
        s
        for slide in prs.slides
        for s in slide.shapes
        if s.has_text_frame and "Line one" in s.text_frame.text
    )
    paragraphs = shape.text_frame.paragraphs
    assert len(paragraphs) == 3
    for p in paragraphs:
        assert p.font.color.rgb == RGBColor.from_string("1A2332")


def test_cli_validate_smoke(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    spec_path = tmp_path / "minimal.json"
    spec_path.write_text(json.dumps(MINIMAL_SPEC), encoding="utf-8")
    exit_code = cli_main(["validate", str(spec_path)])
    captured = capsys.readouterr()
    result = json.loads(captured.out)
    assert exit_code == 0
    assert result["valid"] is True


def test_cli_review_smoke(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    spec_path = tmp_path / "minimal.json"
    spec_path.write_text(json.dumps(MINIMAL_SPEC), encoding="utf-8")
    exit_code = cli_main(["review", str(spec_path)])
    captured = capsys.readouterr()
    result = json.loads(captured.out)
    assert exit_code == 0
    assert "suggestions" in result and "warnings" in result


def test_cli_reference_prints_nonempty_text(
    capsys: pytest.CaptureFixture[str],
) -> None:
    exit_code = cli_main(["reference"])
    captured = capsys.readouterr()
    assert exit_code == 0
    assert len(captured.out) > 500
    assert "render_deck" in captured.out


def test_cli_render_smoke(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    spec_path = tmp_path / "minimal.json"
    spec_path.write_text(json.dumps(MINIMAL_SPEC), encoding="utf-8")
    output = tmp_path / "cli_deck.pptx"
    exit_code = cli_main(["render", str(spec_path), "-o", str(output)])
    captured = capsys.readouterr()
    result = json.loads(captured.out)
    assert exit_code == 0
    assert result["pptx_path"] == str(output)
    assert output.exists()


def test_cli_render_template_flag_overrides_deck_template(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    spec_path = tmp_path / "minimal.json"
    spec_path.write_text(json.dumps(MINIMAL_SPEC), encoding="utf-8")
    output = tmp_path / "cli_deck.pptx"
    exit_code = cli_main(
        ["render", str(spec_path), "-o", str(output), "--template", "modern"]
    )
    assert exit_code == 0
    capsys.readouterr()

    prs = Presentation(str(output))
    fonts = {
        p.font.name
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
        for p in shape.text_frame.paragraphs
    }
    assert fonts == {"Georgia"}


def test_cli_render_unknown_template_flag_fails_cleanly(tmp_path: Path) -> None:
    spec_path = tmp_path / "minimal.json"
    spec_path.write_text(json.dumps(MINIMAL_SPEC), encoding="utf-8")
    exit_code = cli_main(
        [
            "render",
            str(spec_path),
            "-o",
            str(tmp_path / "out.pptx"),
            "--template",
            "nonexistent",
        ]
    )
    assert exit_code == 1


def test_cli_template_extract_writes_a_reviewable_yaml(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    source = tmp_path / "source.pptx"
    Presentation().save(str(source))
    output_dir = tmp_path / "out"

    exit_code = cli_main(
        ["template", "extract", str(source), "acme", "-o", str(output_dir)]
    )
    captured = capsys.readouterr()
    result = json.loads(captured.out)
    assert exit_code == 0
    assert result["template_yaml"] == str(output_dir / "acme.yaml")
    assert (output_dir / "acme.yaml").exists()

    # The written yaml is immediately loadable by Template.from_yaml, same
    # as any hand-authored templates/<name>.yaml.
    loaded = Template.from_yaml(output_dir / "acme.yaml")
    assert loaded.font_family  # never empty, even if the source had no theme font
    assert loaded.font_family == result["font_family"] or result["font_family"] is None


def test_validate_accepts_full_catalog_spec() -> None:
    report = validate(FULL_CATALOG_SPEC)
    assert report.valid is True
    assert report.errors == []


def test_render_deck_writes_all_primitive_types(tmp_path: Path) -> None:

    output = tmp_path / "full_catalog.pptx"
    report = render_deck(FULL_CATALOG_SPEC, output)

    prs = Presentation(str(output))
    assert len(list(prs.slides)) == 3
    shape_types = {shape.shape_type for slide in prs.slides for shape in slide.shapes}

    # Non-negotiable invariant: real, editable shapes only — table and chart
    # are real OOXML graphicFrames, never a flattened picture or video.
    assert MSO_SHAPE_TYPE.PICTURE not in shape_types
    assert MSO_SHAPE_TYPE.TABLE in shape_types
    assert MSO_SHAPE_TYPE.CHART in shape_types

    # image placeholder records a manifest entry instead of a real picture
    assert len(report.manifest) == 1
    assert report.manifest[0]["caption"] == "Team photo goes here"
    assert "rect" in report.manifest[0]


def test_image_placeholder_caption_has_a_visible_font_color(tmp_path: Path) -> None:
    from pptx.dml.color import RGBColor

    output = tmp_path / "full_catalog.pptx"
    render_deck(FULL_CATALOG_SPEC, output)

    prs = Presentation(str(output))
    placeholder = next(
        shape
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame and "Team photo goes here" in shape.text_frame.text
    )
    color = placeholder.text_frame.paragraphs[0].font.color.rgb
    assert color == RGBColor(0x66, 0x66, 0x66)


def test_render_deck_sequence_draws_connectors_between_steps(tmp_path: Path) -> None:

    output = tmp_path / "full_catalog.pptx"
    render_deck(FULL_CATALOG_SPEC, output)

    prs = Presentation(str(output))
    # 4 sequence steps -> 3 connectors between them (plus none from other primitives here)
    connectors = [
        shape
        for slide in prs.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.LINE
    ]
    assert len(connectors) == 3


def test_render_deck_sequence_connectors_run_through_the_gutter_not_the_boxes(
    tmp_path: Path,
) -> None:
    """Connectors must join box edges, never centers — a center-to-center line

    would cross directly over each step's own label text.
    """
    output = tmp_path / "full_catalog.pptx"
    render_deck(FULL_CATALOG_SPEC, output)

    prs = Presentation(str(output))
    step_boxes = [
        shape
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
        and shape.text_frame.text.splitlines()[0]
        in ("Discover", "Design", "Build", "Ship")
    ]
    step_boxes.sort(key=lambda s: s.left)
    assert len(step_boxes) == 4

    connectors = sorted(
        (
            shape
            for slide in prs.slides
            for shape in slide.shapes
            if shape.shape_type.name == "LINE"
        ),
        key=lambda s: s.left,
    )
    assert len(connectors) == 3

    for box, connector in zip(step_boxes, connectors, strict=False):
        box_right_edge = box.left + box.width
        assert connector.left == box_right_edge
        box_vertical_center = box.top + box.height // 2
        assert connector.top == box_vertical_center


# --- Font resolution / chart fonts / per-cell overflow (font+overflow batch) ---


@pytest.mark.parametrize("template_name", ["default", "modern", "classic", "clean"])
def test_resolve_font_path_finds_the_bundled_font_for_every_stock_template(
    template_name: str,
) -> None:
    template = Template.from_name(template_name)
    path = _resolve_font_path(template)
    assert path is not None
    assert path.exists()
    assert path.name == "OpenSans-Regular.ttf"


def test_render_chart_applies_template_font_to_axes(tmp_path: Path) -> None:
    spec = {
        "template": "modern",
        "slides": [
            {
                "header": {"title": "Revenue"},
                "body": [
                    {
                        "primitive": "chart",
                        "chart_type": "bar",
                        "categories": ["Q1", "Q2"],
                        "series": [{"name": "Revenue", "values": [10, 14]}],
                    }
                ],
            }
        ],
    }
    output = tmp_path / "chart.pptx"
    render_deck(spec, output)

    prs = Presentation(str(output))
    chart = next(
        shape.chart for slide in prs.slides for shape in slide.shapes if shape.has_chart
    )
    assert chart.category_axis.tick_labels.font.name == "Georgia"
    assert chart.value_axis.tick_labels.font.name == "Georgia"


def test_render_pie_chart_does_not_crash_applying_axis_font(tmp_path: Path) -> None:
    """Pie charts have neither a category nor a value axis in python-pptx —
    _apply_chart_font must skip them rather than raise.
    """
    spec = {
        "slides": [
            {
                "header": {"title": "Share"},
                "body": [
                    {
                        "primitive": "chart",
                        "chart_type": "pie",
                        "categories": ["A", "B"],
                        "series": [{"name": "Share", "values": [40, 60]}],
                    }
                ],
            }
        ]
    }
    output = tmp_path / "pie.pptx"
    render_deck(spec, output)  # must not raise
    assert output.exists()


def test_table_cell_rects_splits_evenly_into_num_cols_by_num_rows() -> None:
    rect = Rect(x=0, y=0, w=900, h=300)
    cells = _table_cell_rects(rect, num_cols=3, num_rows=3)
    assert len(cells) == 3
    assert all(len(row) == 3 for row in cells)
    assert cells[0][0] == Rect(x=0, y=0, w=300, h=100)
    assert cells[1][2] == Rect(x=600, y=100, w=300, h=100)


def test_table_cell_rects_respects_explicit_col_widths() -> None:
    rect = Rect(x=0, y=0, w=900, h=200)
    cells = _table_cell_rects(rect, num_cols=3, num_rows=2, col_widths=[100, 300, 500])
    assert cells[0][0] == Rect(x=0, y=0, w=100, h=100)
    assert cells[0][1] == Rect(x=100, y=0, w=300, h=100)
    assert cells[0][2] == Rect(x=400, y=0, w=500, h=100)
    assert cells[1][2] == Rect(x=400, y=100, w=500, h=100)


@pytest.fixture
def monospace_metrics():
    """A trivial monospace stand-in — every character advances the same
    width — so column-width proportions are exercised without depending on
    the bundled font's real glyph metrics.
    """
    from compono.validator import FontMetrics

    return FontMetrics(units_per_em=1000, default_advance=600, advance_widths={})


def test_table_column_widths_emu_sums_to_rect_width(monospace_metrics) -> None:
    from compono.render import _table_column_widths_emu

    rect = Rect(x=0, y=0, w=900_000, h=200_000)
    widths = _table_column_widths_emu(
        headers=["A", "A much longer header here", "B"],
        rows=[["x", "y", "z"]],
        rect=rect,
        font_metrics=monospace_metrics,
    )
    assert widths is not None
    assert sum(widths) == rect.w
    assert len(widths) == 3


def test_table_column_widths_emu_gives_more_space_to_the_longer_column(
    monospace_metrics,
) -> None:
    from compono.render import _table_column_widths_emu

    rect = Rect(x=0, y=0, w=900_000, h=200_000)
    widths = _table_column_widths_emu(
        headers=["Short", "This is a considerably longer column header"],
        rows=[["s", "l"]],
        rect=rect,
        font_metrics=monospace_metrics,
    )
    assert widths is not None
    assert widths[1] > widths[0]


def test_table_column_widths_emu_returns_none_without_font_metrics() -> None:
    from compono.render import _table_column_widths_emu

    rect = Rect(x=0, y=0, w=900_000, h=200_000)
    widths = _table_column_widths_emu(
        headers=["A", "B"], rows=[["1", "2"]], rect=rect, font_metrics=None
    )
    assert widths is None


def test_render_table_sets_proportional_column_widths_when_font_available(
    tmp_path: Path,
) -> None:
    spec = {
        "slides": [
            {
                "body": [
                    {
                        "primitive": "table",
                        "headers": ["ID", "A considerably longer description column"],
                        "rows": [["1", "short"]],
                    }
                ]
            }
        ]
    }
    output = tmp_path / "deck.pptx"
    render_deck(spec, output, template=Template.from_yaml())

    prs = Presentation(str(output))
    tbl = next(s.table for s in next(iter(prs.slides)).shapes if s.has_table)
    widths = [col.width for col in tbl.columns]
    assert widths[1] > widths[0]


def test_sequence_step_rects_splits_evenly_left_to_right() -> None:
    rect = Rect(x=0, y=0, w=400, h=100)
    steps = _sequence_step_rects(rect, num_steps=4)
    assert steps == [
        Rect(x=0, y=0, w=100, h=100),
        Rect(x=100, y=0, w=100, h=100),
        Rect(x=200, y=0, w=100, h=100),
        Rect(x=300, y=0, w=100, h=100),
    ]


def test_extract_text_fields_checks_each_table_cell_against_its_own_sub_rect() -> None:
    """The old combined-text heuristic checked all headers+rows joined
    against the whole rect — a single overlong cell could be masked by
    plenty of short cells nearby. Per-cell checking must not do that.
    """
    table = Table(headers=["A", "B"], rows=[["short", "short"]])
    rect = Rect(x=0, y=0, w=1000, h=500)
    fields = _extract_text_fields(table, rect)

    field_names = {name for name, _, _, _ in fields}
    assert field_names == {"headers[0]", "headers[1]", "rows[0][0]", "rows[0][1]"}

    for _, _, _, sub_rect in fields:
        assert sub_rect is not None
        assert sub_rect.w == 500  # 2 cols
        assert sub_rect.h == 250  # header row + 1 data row


def test_extract_text_fields_checks_each_sequence_step_against_its_own_sub_rect() -> (
    None
):
    sequence = Sequence(
        steps=[
            SequenceStep(label="One"),
            SequenceStep(label="Two"),
        ],
        orientation="horizontal",
    )
    rect = Rect(x=0, y=0, w=1000, h=500)
    fields = _extract_text_fields(sequence, rect)

    assert [name for name, _, _, _ in fields] == ["steps[0]", "steps[1]"]
    for _, _, _, sub_rect in fields:
        assert sub_rect is not None
        assert sub_rect.w == 500
        assert sub_rect.h == 500


def test_validate_flags_a_single_overlong_table_cell_even_though_the_combined_text_would_fit() -> (
    None
):
    """Regression test for the actual bug the per-cell fix addresses: many
    short cells plus one very long cell used to pass because only the
    combined string was checked against the whole table's rect.
    """
    long_cell = "word " * 400  # long enough to overflow a single narrow cell
    spec = {
        "slides": [
            {
                "header": {"title": "Comparison"},
                "body": [
                    {
                        "primitive": "table",
                        "headers": ["A", "B", "C", "D"],
                        "rows": [["x", "y", long_cell, "z"]],
                    }
                ],
            }
        ]
    }
    report = validate(spec)
    assert report.valid is False
    assert any(
        e["error"] == "overflow" and e["field"] == "rows[0][2]" for e in report.errors
    )


def test_render_deck_diagram_renders_nodes_and_edges_as_real_shapes(
    tmp_path: Path,
) -> None:

    spec = {
        "slides": [
            {
                "header": {"title": "Pipeline"},
                "body": [
                    {
                        "primitive": "diagram",
                        "nodes": [
                            {"label": "User"},
                            {"label": "Router"},
                            {"label": "Retriever", "fill": "#FF0000"},
                        ],
                    }
                ],
            }
        ]
    }
    output = tmp_path / "diagram.pptx"
    render_deck(spec, output)

    prs = Presentation(str(output))
    slide = next(iter(prs.slides))
    shape_types = [s.shape_type for s in slide.shapes]

    # Non-negotiable invariant: real, editable shapes only.
    assert MSO_SHAPE_TYPE.PICTURE not in shape_types
    # header + 3 nodes + 2 connectors (default linear chain) + footer = 7.
    assert len(slide.shapes) == 7

    node_texts = {
        s.text_frame.text
        for s in slide.shapes
        if s.has_text_frame and s.text_frame.text
    }
    assert {"User", "Router", "Retriever"} <= node_texts

    # A node's own fill override wins over the diagram-level default.
    retriever = next(
        s for s in slide.shapes if s.has_text_frame and s.text_frame.text == "Retriever"
    )
    assert retriever.fill.fore_color.rgb == RGBColor.from_string("FF0000")


# --- Template branding: primary_color/accent_color/logo_path (all additive/gated) ---


def _table_spec() -> dict:
    return {
        "slides": [
            {
                "body": [
                    {
                        "primitive": "table",
                        "headers": ["Metric", "Value"],
                        "rows": [["Revenue", "$1.2M"]],
                    }
                ]
            }
        ]
    }


def test_table_header_fill_unset_by_default(tmp_path: Path) -> None:
    output = tmp_path / "deck.pptx"
    render_deck(_table_spec(), output, template=Template.from_yaml())

    prs = Presentation(str(output))
    tbl = next(s.table for s in next(iter(prs.slides)).shapes if s.has_table)
    # No template.primary_color set -> python-pptx's own default cell fill,
    # never an explicitly-set solid color.
    assert tbl.cell(0, 0).fill.type is None


def test_table_header_fill_applies_primary_color_when_set(tmp_path: Path) -> None:
    branded = dataclasses.replace(Template.from_yaml(), primary_color="#1F4E79")
    output = tmp_path / "deck.pptx"
    render_deck(_table_spec(), output, template=branded)

    prs = Presentation(str(output))
    tbl = next(s.table for s in next(iter(prs.slides)).shapes if s.has_table)
    assert tbl.cell(0, 0).fill.fore_color.rgb == RGBColor.from_string("1F4E79")
    assert tbl.cell(0, 0).text_frame.paragraphs[0].font.color.rgb == RGBColor(
        0xFF, 0xFF, 0xFF
    )


def _sequence_spec() -> dict:
    return {
        "slides": [
            {
                "body": [
                    {
                        "primitive": "sequence",
                        "steps": [{"label": "Discovery"}, {"label": "Delivery"}],
                    }
                ]
            }
        ]
    }


def test_sequence_step_fill_applies_accent_color_when_set(tmp_path: Path) -> None:
    branded = dataclasses.replace(Template.from_yaml(), accent_color="#2E86AB")
    output = tmp_path / "deck.pptx"
    render_deck(_sequence_spec(), output, template=branded)

    prs = Presentation(str(output))
    step_shapes = [
        s
        for s in next(iter(prs.slides)).shapes
        if s.has_text_frame and s.text_frame.text == "Discovery"
    ]
    assert len(step_shapes) == 1
    assert step_shapes[0].fill.fore_color.rgb == RGBColor.from_string("2E86AB")


def test_header_logo_renders_a_real_picture_when_logo_path_set(tmp_path: Path) -> None:
    from PIL import Image as PILImage

    logo_path = tmp_path / "logo.png"
    PILImage.new("RGB", (40, 20), color=(10, 20, 30)).save(logo_path)
    branded = dataclasses.replace(Template.from_yaml(), logo_path=logo_path)

    output = tmp_path / "deck.pptx"
    render_deck(
        {"slides": [{"header": {"title": "Q3 Results"}, "body": []}]},
        output,
        template=branded,
    )

    prs = Presentation(str(output))
    pictures = [
        s
        for s in next(iter(prs.slides)).shapes
        if s.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert len(pictures) == 1


def test_header_has_no_picture_when_logo_path_unset(tmp_path: Path) -> None:
    output = tmp_path / "deck.pptx"
    render_deck(
        {"slides": [{"header": {"title": "Q3 Results"}, "body": []}]},
        output,
        template=Template.from_yaml(),
    )

    prs = Presentation(str(output))
    pictures = [
        s
        for s in next(iter(prs.slides)).shapes
        if s.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert pictures == []


# --- Table cell_fills + gantt primitive ---


def test_render_table_applies_cell_fills(tmp_path: Path) -> None:
    spec = {
        "slides": [
            {
                "body": [
                    {
                        "primitive": "table",
                        "headers": ["A", "B"],
                        "rows": [["1", "2"], ["3", "4"]],
                        "cell_fills": [{"row": 1, "col": 0, "fill": "#2A6FDB"}],
                    }
                ]
            }
        ]
    }
    output = tmp_path / "deck.pptx"
    render_deck(spec, output, template=Template.from_yaml())

    prs = Presentation(str(output))
    tbl = next(s.table for s in next(iter(prs.slides)).shapes if s.has_table)
    # Body row 1 (0-based, "3") -> pptx row index 2 (header is row 0).
    assert tbl.cell(2, 0).fill.fore_color.rgb == RGBColor.from_string("2A6FDB")
    # Every other body cell stays unfilled.
    assert tbl.cell(1, 0).fill.type is None
    assert tbl.cell(1, 1).fill.type is None


def test_render_table_merges_cells_and_clears_the_non_origin_cells_text(
    tmp_path: Path,
) -> None:
    spec = {
        "slides": [
            {
                "body": [
                    {
                        "primitive": "table",
                        "headers": ["Region", "Q1", "Q2"],
                        "rows": [
                            ["North", "10", "12"],
                            ["North", "11", "13"],
                            ["South", "5", "6"],
                        ],
                        "merges": [{"row1": 0, "col1": 0, "row2": 1, "col2": 0}],
                    }
                ]
            }
        ]
    }
    output = tmp_path / "deck.pptx"
    render_deck(spec, output, template=Template.from_yaml())

    prs = Presentation(str(output))
    tbl = next(s.table for s in next(iter(prs.slides)).shapes if s.has_table)
    # Merged region is pptx rows 1-2, col 0 (header is row 0).
    assert tbl.cell(1, 0).is_merge_origin
    # The merge must not concatenate both original "North" values.
    assert tbl.cell(1, 0).text == "North"
    assert tbl.cell(2, 0).text == ""
    # Unmerged row is untouched.
    assert tbl.cell(3, 0).text == "South"
    assert tbl.cell(2, 1).fill.type is None


def test_render_gantt_produces_a_real_table_with_correct_task_spans(
    tmp_path: Path,
) -> None:
    spec = {
        "slides": [
            {
                "header": {"title": "Project Timeline"},
                "body": [
                    {
                        "primitive": "gantt",
                        "unit_labels": ["Wk 1", "Wk 2", "Wk 3", "Wk 4"],
                        "task_fill": "#2A6FDB",
                        "tasks": [
                            {
                                "label": "Discovery",
                                "start_unit": 0,
                                "duration_units": 2,
                            },
                            {
                                "label": "Design",
                                "start_unit": 2,
                                "duration_units": 2,
                                "fill": "#D9534F",
                            },
                        ],
                    }
                ],
            }
        ]
    }
    output = tmp_path / "gantt.pptx"
    render_deck(spec, output)

    prs = Presentation(str(output))
    slide = next(iter(prs.slides))
    # A real, editable table — never a flattened image.
    assert MSO_SHAPE_TYPE.PICTURE not in [s.shape_type for s in slide.shapes]
    tbl = next(s.table for s in slide.shapes if s.has_table)

    assert [tbl.cell(0, c).text for c in range(5)] == [
        "Task",
        "Wk 1",
        "Wk 2",
        "Wk 3",
        "Wk 4",
    ]
    assert tbl.cell(1, 0).text == "Discovery"
    assert tbl.cell(2, 0).text == "Design"

    # Discovery spans Wk 1-2 (cols 1-2) with the default task_fill.
    assert tbl.cell(1, 1).fill.fore_color.rgb == RGBColor.from_string("2A6FDB")
    assert tbl.cell(1, 2).fill.fore_color.rgb == RGBColor.from_string("2A6FDB")
    assert tbl.cell(1, 3).fill.type is None
    # Design spans Wk 3-4 (cols 3-4) with its own override color.
    assert tbl.cell(2, 3).fill.fore_color.rgb == RGBColor.from_string("D9534F")
    assert tbl.cell(2, 4).fill.fore_color.rgb == RGBColor.from_string("D9534F")
    assert tbl.cell(2, 1).fill.type is None
