"""Unit tests for src/compono/inspire.py — scan/aggregate/write_skill.

Fixtures render tiny specs through compono's own render_deck() into real
.pptx files (no binary fixtures committed), then scan those.
"""

import json
from pathlib import Path

import pytest

from compono.cli import main as cli_main
from compono.inspire import aggregate, scan_deck, write_skill
from compono.render import render_deck


def _grid_spec(columns: int, fill: str) -> dict:
    return {
        "slides": [
            {
                "header": {"title": "Grid deck"},
                "body": [
                    {
                        "primitive": "grid",
                        "columns": columns,
                        "items": [
                            {
                                "primitive": "shape",
                                "kind": "rect",
                                "fill": fill,
                                "text": {"content": f"Card {i}", "color": "#FFFFFF"},
                            }
                            for i in range(columns)
                        ],
                    }
                ],
            }
        ]
    }


@pytest.fixture
def three_col_deck(tmp_path: Path) -> Path:
    out = tmp_path / "three_col.pptx"
    render_deck(_grid_spec(3, "#2A9D8F"), str(out))
    return out


@pytest.fixture
def five_col_deck(tmp_path: Path) -> Path:
    out = tmp_path / "five_col.pptx"
    render_deck(_grid_spec(5, "#2A9D8F"), str(out))
    return out


@pytest.fixture
def one_off_color_deck(tmp_path: Path) -> Path:
    out = tmp_path / "one_off.pptx"
    render_deck(_grid_spec(3, "#F4A261"), str(out))
    return out


def test_scan_deck_never_returns_literal_source_text(three_col_deck: Path) -> None:
    profile = scan_deck(three_col_deck)
    dumped = json.dumps(profile)
    # None of the literal shape.text.content strings we rendered should ever
    # appear in the profile — only structural facts (colors, fonts, counts).
    assert "Card 0" not in dumped
    assert "Grid deck" not in dumped


def test_scan_deck_recovers_known_grid_column_count(three_col_deck: Path) -> None:
    profile = scan_deck(three_col_deck)
    columns_seen = {g["columns"] for g in profile["grids_detected"]}
    assert 3 in columns_seen


def test_scan_deck_detected_grids_carry_a_confidence_score(
    three_col_deck: Path,
) -> None:
    profile = scan_deck(three_col_deck)
    assert profile["grids_detected"], "expected at least one detected grid"
    for g in profile["grids_detected"]:
        assert 0.0 <= g["confidence"] <= 1.0


def test_scan_deck_omits_low_confidence_rows_rather_than_guessing(
    tmp_path: Path,
) -> None:
    # A slide with no shapes forming a real row-aligned pattern (a single
    # header-only slide) must never fabricate a grid.
    spec = {"slides": [{"header": {"title": "Just a title"}}]}
    out = tmp_path / "lone.pptx"
    render_deck(spec, str(out))
    profile = scan_deck(out)
    assert profile["grids_detected"] == []


def test_aggregate_promotes_color_seen_in_majority_of_decks(
    three_col_deck: Path, five_col_deck: Path, one_off_color_deck: Path
) -> None:
    profile = aggregate(
        [three_col_deck, five_col_deck, one_off_color_deck], min_repeat_ratio=0.5
    )
    assert "2A9D8F" in profile["recurring_palette"]
    assert "F4A261" in profile["one_off_palette"]
    assert "F4A261" not in profile["recurring_palette"]


def test_aggregate_does_not_let_a_many_slide_deck_dominate(tmp_path: Path) -> None:
    # A 10-slide deck all using one color vs. three 1-slide decks using
    # another color each — the many-slide deck's color must not be counted
    # more than once (it's one deck, not ten votes).
    many_slides = {
        "slides": [
            {
                "header": {"title": f"Slide {i}"},
                "body": [
                    {
                        "primitive": "shape",
                        "kind": "rect",
                        "fill": "#111111",
                    }
                ],
            }
            for i in range(10)
        ]
    }
    many_path = tmp_path / "many.pptx"
    render_deck(many_slides, str(many_path))

    other_paths = []
    for i, color in enumerate(["#222222", "#333333", "#444444"]):
        spec = {
            "slides": [
                {
                    "header": {"title": "One slide"},
                    "body": [{"primitive": "shape", "kind": "rect", "fill": color}],
                }
            ]
        }
        p = tmp_path / f"other_{i}.pptx"
        render_deck(spec, str(p))
        other_paths.append(p)

    profile = aggregate([many_path, *other_paths], min_repeat_ratio=0.5)
    # #111111 appears in exactly 1 of 4 decks, same as each of the others —
    # none should be "recurring" at a 50% threshold.
    assert profile["recurring_palette"] == []


def test_write_skill_produces_both_files_with_valid_frontmatter(
    tmp_path: Path, three_col_deck: Path
) -> None:
    profile = aggregate([three_col_deck])
    out_dir = tmp_path / "skills" / "inspire-mystyle"
    files = write_skill(profile, out_dir, name="mystyle")

    assert files.skill_md.exists()
    assert files.profile_json.exists()

    skill_text = files.skill_md.read_text(encoding="utf-8")
    assert skill_text.startswith("---\n")
    assert "name: inspire-mystyle" in skill_text

    round_tripped = json.loads(files.profile_json.read_text(encoding="utf-8"))
    assert round_tripped["n_example_decks"] == 1


def test_aggregate_excludes_unset_font_from_recurring_fonts(tmp_path: Path) -> None:
    # A run with no explicit font override (font.name is None -> scan_deck
    # reports family "?") must never crowd out, or itself be reported as, a
    # recurring font — it carries no actionable signal about what font a
    # deck actually renders in. compono itself always sets an explicit
    # font, so this strips it back off after rendering to simulate a
    # hand-authored deck that relies on inherited theme fonts.
    from pptx import Presentation

    spec = {
        "slides": [
            {
                "header": {"title": "No explicit font"},
                "body": [{"primitive": "text", "content": ["Just some text"]}],
            }
        ]
    }
    out = tmp_path / "no_font.pptx"
    render_deck(spec, str(out))

    prs = Presentation(str(out))
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.name = None
    prs.save(str(out))

    profile = aggregate([out, out])
    assert all(f["family"] != "?" for f in profile["recurring_fonts"])


def test_aggregate_raises_on_empty_path_list() -> None:
    with pytest.raises(ValueError, match="at least one|no scannable"):
        aggregate([])


def test_cli_inspire_scan_smoke(tmp_path: Path, three_col_deck: Path) -> None:
    decks_dir = tmp_path / "decks"
    decks_dir.mkdir()
    (decks_dir / "deck.pptx").write_bytes(three_col_deck.read_bytes())

    out_dir = tmp_path / "out"
    exit_code = cli_main(
        ["inspire", "scan", str(decks_dir), "-o", str(out_dir), "--name", "team"]
    )
    assert exit_code == 0
    assert (out_dir / "SKILL.md").exists()
    assert (out_dir / "profile.json").exists()
