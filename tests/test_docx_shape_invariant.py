"""Golden/invariant test: every docx primitive except `chart` produces a
genuine, editable python-docx object — never a flattened image.

`chart` is the one documented exception (docx_schema.DocChart's docstring):
`python-docx` has no native chart API, so it is deliberately rasterized via
matplotlib and embedded as a picture. This test locks that exception down to
exactly one primitive, so it can't silently spread to others.
"""

from pathlib import Path

import docx as python_docx

from compono.docx import render_docx

SPEC = {
    "title": "Shape Invariant Check",
    "sections": [
        {
            "body": [
                {"primitive": "heading", "text": "Title", "level": 1},
                {"primitive": "paragraph", "runs": [{"text": "Body text."}]},
                {"primitive": "bullet_list", "items": [[{"text": "Bullet"}]]},
                {"primitive": "numbered_list", "items": [[{"text": "Numbered"}]]},
                {"primitive": "table", "headers": ["A"], "rows": [["1"]]},
                {"primitive": "image", "placeholder": True, "caption": "Logo"},
                {"primitive": "page_break"},
                {
                    "primitive": "chart",
                    "chart_type": "line",
                    "categories": ["A", "B"],
                    "series": [{"name": "s", "values": [1, 2]}],
                },
            ]
        }
    ],
}


def test_only_the_chart_primitive_renders_as_a_picture(tmp_path: Path) -> None:
    output = tmp_path / "invariant.docx"
    render_docx(SPEC, output)

    doc = python_docx.Document(str(output))

    # Real python-docx objects for every non-chart primitive.
    assert doc.tables, "table primitive did not produce a real python-docx table"
    heading = next(p for p in doc.paragraphs if p.text == "Title")
    assert heading.style.name.startswith("Heading")
    assert any(p.style.name == "List Bullet" for p in doc.paragraphs)
    assert any(p.style.name == "List Number" for p in doc.paragraphs)

    # Exactly one picture in the whole document: the chart. The image
    # primitive here is a placeholder (text only), so it contributes none.
    assert len(doc.inline_shapes) == 1
