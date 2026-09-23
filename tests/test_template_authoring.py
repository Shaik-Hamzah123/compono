"""Unit tests for src/compono/template_authoring.py (`compono template extract`).

Fixtures build a real .pptx with python-pptx rather than hand-editing OOXML,
except for the one thing python-pptx's write API won't do itself — placing a
picture on the slide master — which is constructed via the same low-level
`CT_Picture.new_pic` + `get_or_add_image_part` primitives python-pptx itself
uses internally for a normal slide picture.
"""

import io
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml.shapes.picture import CT_Picture
from pptx.util import Emu

from compono.resolver import Template
from compono.template_authoring import extract_template_data, write_extracted_template


def _make_logo_bytes() -> bytes:
    from PIL import Image as PILImage

    buf = io.BytesIO()
    PILImage.new("RGB", (40, 20), color=(10, 20, 30)).save(buf, format="PNG")
    return buf.getvalue()


def _build_fixture_pptx(path: Path, *, with_logo: bool) -> None:
    prs = Presentation()
    if with_logo:
        master = prs.slide_masters[0]
        _image_part, rId = master.part.get_or_add_image_part(
            io.BytesIO(_make_logo_bytes())
        )
        pic = CT_Picture.new_pic(
            9001, "Logo", "logo", rId, Emu(0), Emu(0), Emu(500000), Emu(300000)
        )
        master.shapes.element.append(pic)
    prs.save(str(path))


@pytest.fixture
def fixture_pptx(tmp_path: Path) -> Path:
    path = tmp_path / "source.pptx"
    _build_fixture_pptx(path, with_logo=True)
    return path


def test_extract_template_data_reads_real_page_size(fixture_pptx: Path) -> None:
    data = extract_template_data(fixture_pptx)
    prs = Presentation(str(fixture_pptx))
    assert data.page_width_in == pytest.approx(prs.slide_width / 914400)
    assert data.page_height_in == pytest.approx(prs.slide_height / 914400)


def test_extract_template_data_finds_font_and_colors_from_the_real_default_theme(
    fixture_pptx: Path,
) -> None:
    """A fresh python-pptx Presentation() always ships a real theme part —
    assert extraction is non-None and well-formed, not a specific hex value
    (python-pptx's bundled default theme is an implementation detail, not
    something this repo documents or should pin against).
    """
    data = extract_template_data(fixture_pptx)
    assert data.font_family
    assert isinstance(data.font_family, str)
    for color in (data.primary_color, data.accent_color):
        assert color is not None
        assert color.startswith("#")
        assert len(color) == 7
        int(color[1:], 16)  # valid hex


def test_extract_template_data_finds_the_master_logo(fixture_pptx: Path) -> None:
    data = extract_template_data(fixture_pptx)
    assert data.logo_bytes == _make_logo_bytes()
    assert data.logo_ext == "png"


def test_extract_template_data_logo_is_none_without_one(tmp_path: Path) -> None:
    path = tmp_path / "no_logo.pptx"
    _build_fixture_pptx(path, with_logo=False)
    data = extract_template_data(path)
    assert data.logo_bytes is None
    assert data.logo_ext is None


def test_extract_template_data_never_raises_on_missing_theme_part(
    tmp_path: Path,
) -> None:
    """A malformed/theme-less .pptx (simulated by pointing at a file whose
    zip has no ppt/theme/theme1.xml) must degrade to None fields, not raise.
    """
    import zipfile

    path = tmp_path / "no_theme.pptx"
    _build_fixture_pptx(path, with_logo=False)

    # Strip the theme part out of an otherwise-valid pptx zip.
    stripped = tmp_path / "stripped.pptx"
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(stripped, "w") as dst:
        for item in src.infolist():
            if item.filename == "ppt/theme/theme1.xml":
                continue
            dst.writestr(item, src.read(item.filename))

    data = extract_template_data(stripped)
    assert data.font_family is None
    assert data.primary_color is None
    assert data.accent_color is None


def test_write_extracted_template_round_trips_through_template_from_yaml(
    fixture_pptx: Path, tmp_path: Path
) -> None:
    data = extract_template_data(fixture_pptx)
    output_dir = tmp_path / "out"
    yaml_path = write_extracted_template(data, "acme", output_dir)

    assert yaml_path == output_dir / "acme.yaml"
    assert yaml_path.exists()

    template = Template.from_yaml(yaml_path)
    assert template.page_width == round(data.page_width_in * 914400)
    assert template.font_family == data.font_family
    assert template.primary_color == data.primary_color
    assert template.accent_color == data.accent_color
    assert template.logo_path is not None
    assert template.logo_path.read_bytes() == data.logo_bytes

    # Margins/header/footer/gutter are copied from compono's own default —
    # never derived from the source deck.
    default = Template.from_yaml()
    assert template.margin_top == default.margin_top
    assert template.header_height == default.header_height
    assert template.footer_height == default.footer_height
    assert template.gutter == default.gutter


def test_write_extracted_template_omits_colors_and_logo_keys_when_absent(
    tmp_path: Path,
) -> None:
    path = tmp_path / "no_logo.pptx"
    _build_fixture_pptx(path, with_logo=False)
    data = extract_template_data(path)
    # Force colors to None too, independent of whatever theme values the
    # fixture happened to have, to exercise the "nothing found" path cleanly.
    import dataclasses

    data = dataclasses.replace(data, primary_color=None, accent_color=None)

    output_dir = tmp_path / "out"
    yaml_path = write_extracted_template(data, "bare", output_dir)
    template = Template.from_yaml(yaml_path)
    assert template.primary_color is None
    assert template.accent_color is None
    assert template.logo_path is None
    assert not (output_dir / "assets").exists()
