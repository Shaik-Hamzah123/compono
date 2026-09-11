"""Orchestrates validate -> resolve -> write pptx.

Exposes the two public verbs (COMPONO_PLAN.md sections 8-9): render_deck(spec)
and validate(spec). No client objects, no session lifecycle — both are pure
functions over a Deck spec (raw dict or a typed Deck).

Full v1 primitive catalog: header, text, image, stat, grid, table, sequence,
chart, shape (COMPONO_PLAN.md section 5).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from itertools import pairwise
from pathlib import Path
from typing import Any

from PIL import Image as PILImage
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt
from pydantic import ValidationError
from pydantic_core import ErrorDetails

from compono.resolver import (
    ConnectorPoints,
    LayoutResult,
    Rect,
    Template,
    resolve_slide,
)
from compono.schema import (
    Chart,
    Deck,
    Header,
    Image,
    PrimitiveBase,
    Sequence,
    Shape,
    Slide,
    Stat,
    Table,
    Text,
)
from compono.validator import (
    SAFE_FONTS,
    FontMetrics,
    build_overflow_error,
    check_overflow,
    load_font_metrics,
    resolve_safe_font,
)

HEADER_FONT_SIZE_PT = 28
BODY_FONT_SIZE_PT = 18
STAT_VALUE_FONT_SIZE_PT = 36
STAT_LABEL_FONT_SIZE_PT = 14
TABLE_FONT_SIZE_PT = 14
CAPTION_FONT_SIZE_PT = 12

# Fallback system fonts used only until a real font ships in src/compono/fonts/
# (SAFE_FONTS in validator.py is still empty). Overflow validation is skipped,
# not faked, when none of these are found either.
_FALLBACK_SYSTEM_FONTS = [
    Path("C:/Windows/Fonts/arial.ttf"),
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
]

_SHAPE_KIND_TO_MSO = {
    "rect": MSO_SHAPE.RECTANGLE,
    "rounded_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "oval": MSO_SHAPE.OVAL,
}
_ALIGN_TO_PP = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
}
_VALIGN_TO_MSO = {
    "top": MSO_ANCHOR.TOP,
    "middle": MSO_ANCHOR.MIDDLE,
    "bottom": MSO_ANCHOR.BOTTOM,
}
_CHART_TYPE_TO_XL = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
}


class DeckValidationError(Exception):
    """The one exception type. Carries the structured list of per-field errors."""

    def __init__(self, errors: list[dict[str, Any]]) -> None:
        super().__init__(f"{len(errors)} validation error(s)")
        self.errors = errors


@dataclass
class ValidationReport:
    valid: bool
    errors: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass
class RenderReport:
    """render_deck returns this, not just a file — a feedback channel (COMPONO_PLAN.md section 8, item 5)."""

    pptx_path: Path
    manifest: list[dict[str, Any]]
    warnings: list[str]
    actual_layout: list[LayoutResult]


def _resolve_font_path() -> Path | None:
    for name in SAFE_FONTS:
        path = resolve_safe_font(name)
        if path is not None and path.exists():
            return path
    return next((p for p in _FALLBACK_SYSTEM_FONTS if p.exists()), None)


def _pydantic_error_to_dict(err: ErrorDetails) -> dict[str, Any]:
    loc = ".".join(str(part) for part in err["loc"]) or "<deck>"
    return {
        "slide": None,
        "primitive": loc,
        "field": err["loc"][-1] if err["loc"] else None,
        "error": err["type"],
        "detail": err["msg"],
        "fix": "Check the field against the schema description and correct the value/type.",
    }


def _parse_deck(spec: dict[str, Any] | Deck) -> Deck:
    if isinstance(spec, Deck):
        return spec
    try:
        return Deck.model_validate(spec)
    except ValidationError as exc:
        raise DeckValidationError(
            [_pydantic_error_to_dict(e) for e in exc.errors()]
        ) from exc


def _extract_text_fields(
    primitive: PrimitiveBase | None,
) -> list[tuple[str, str, float]]:
    """(field_name, text, font_size_pt) for every text-bearing field on a primitive.

    Every entry is checked through the same shared check_overflow/wrap_lines
    routine (validator.py) — primitives never grow their own wrap logic
    (COMPONO_PLAN.md section 5, "Text-in-shape").
    """
    if isinstance(primitive, Header):
        return [("title", primitive.title, HEADER_FONT_SIZE_PT)]
    if isinstance(primitive, Text):
        content = (
            primitive.content
            if isinstance(primitive.content, str)
            else " ".join(primitive.content)
        )
        return [("content", content, BODY_FONT_SIZE_PT)]
    if isinstance(primitive, Shape) and primitive.text is not None:
        return [("text.content", primitive.text.content, BODY_FONT_SIZE_PT)]
    if isinstance(primitive, Stat):
        return [
            ("value", primitive.value, STAT_VALUE_FONT_SIZE_PT),
            ("label", primitive.label, STAT_LABEL_FONT_SIZE_PT),
        ]
    if isinstance(primitive, Table):
        # A rough combined-text check for v1 — not per-cell overflow yet.
        combined = (
            " ".join(primitive.headers)
            + " "
            + " ".join(" ".join(row) for row in primitive.rows)
        )
        return [("rows", combined, TABLE_FONT_SIZE_PT)]
    if isinstance(primitive, Sequence):
        combined = " ".join(
            f"{step.label}: {step.description}" if step.description else step.label
            for step in primitive.steps
        )
        return [("steps", combined, BODY_FONT_SIZE_PT)]
    if isinstance(primitive, Image) and primitive.caption:
        return [("caption", primitive.caption, CAPTION_FONT_SIZE_PT)]
    return []


def _check_slide(
    slide_index: int,
    slide: Slide,
    template: Template,
    font_metrics: FontMetrics | None,
) -> tuple[LayoutResult | None, list[dict[str, Any]], list[str]]:
    errors: list[dict[str, Any]] = []
    warnings: list[str] = []

    try:
        layout = resolve_slide(template, header=slide.header, body=slide.body)
    except ValueError as exc:
        errors.append(
            {
                "slide": slide_index,
                "primitive": "slide",
                "field": None,
                "error": "layout",
                "detail": str(exc),
                "fix": "Adjust the template margins/header/footer heights, or reduce body item count.",
            }
        )
        return None, errors, warnings

    if font_metrics is None:
        warnings.append(
            f"slide {slide_index}: no font available — overflow validation skipped."
        )
        return layout, errors, warnings

    for item_id, rect in layout.rects.items():
        box_width_pt = Emu(rect.w).pt
        box_height_pt = Emu(rect.h).pt
        for field_name, text, font_size_pt in _extract_text_fields(
            layout.items.get(item_id)
        ):
            report = check_overflow(
                text, font_metrics, font_size_pt, box_width_pt, box_height_pt
            )
            if report.overflow:
                errors.append(
                    build_overflow_error(
                        slide_index, item_id, field_name, report, font_size_pt
                    )
                )

    return layout, errors, warnings


def _unknown_template_error(detail: str) -> dict[str, Any]:
    return {
        "slide": None,
        "primitive": "template",
        "field": "template",
        "error": "unknown_template",
        "detail": detail,
        "fix": "Use one of the available template names listed above, or omit `template` for the default.",
    }


def resolve_deck_template(deck: Deck, template: Template | None) -> Template:
    """An explicit `template=` kwarg always wins; otherwise resolve the name the
    spec itself asked for (`Deck.template`, schema.py) — previously unwired,
    every render silently used the hardcoded default regardless of this field.
    Raises ValueError (from Template.from_name) on an unknown name — shared by
    validate/render_deck/review, each of which reports it in their own shape.
    """
    if template is not None:
        return template
    return Template.from_name(deck.template)


def validate(
    spec: dict[str, Any] | Deck, *, template: Template | None = None
) -> ValidationReport:
    """Cheap and separate from render — no pptx write, no image render (COMPONO_PLAN.md section 8, item 4)."""
    try:
        deck = _parse_deck(spec)
    except DeckValidationError as exc:
        return ValidationReport(valid=False, errors=exc.errors)

    try:
        resolved_template = resolve_deck_template(deck, template)
    except ValueError as exc:
        return ValidationReport(valid=False, errors=[_unknown_template_error(str(exc))])

    font_path = _resolve_font_path()
    font_metrics = load_font_metrics(font_path) if font_path is not None else None

    all_errors: list[dict[str, Any]] = []
    all_warnings: list[str] = []
    for i, slide in enumerate(deck.slides):
        _, errors, warnings = _check_slide(i, slide, resolved_template, font_metrics)
        all_errors.extend(errors)
        all_warnings.extend(warnings)

    return ValidationReport(
        valid=not all_errors, errors=all_errors, warnings=all_warnings
    )


def render_deck(
    spec: dict[str, Any] | Deck,
    output_path: str | Path,
    *,
    template: Template | None = None,
) -> RenderReport:
    output_path = Path(output_path)

    deck = _parse_deck(spec)

    try:
        resolved_template = resolve_deck_template(deck, template)
    except ValueError as exc:
        raise DeckValidationError([_unknown_template_error(str(exc))]) from exc

    font_path = _resolve_font_path()
    font_metrics = load_font_metrics(font_path) if font_path is not None else None

    layouts: list[LayoutResult] = []
    errors: list[dict[str, Any]] = []
    warnings: list[str] = []

    for i, slide in enumerate(deck.slides):
        layout, slide_errors, slide_warnings = _check_slide(
            i, slide, resolved_template, font_metrics
        )
        errors.extend(slide_errors)
        warnings.extend(slide_warnings)
        layouts.append(layout)  # type: ignore[arg-type]  # None only paired with an error, checked below

    if errors:
        raise DeckValidationError(errors)

    prs = Presentation()
    prs.slide_width = Emu(resolved_template.page_width)
    prs.slide_height = Emu(resolved_template.page_height)
    blank_layout = prs.slide_layouts[6]

    manifest: list[dict[str, Any]] = []
    for i, layout in enumerate(layouts):
        pptx_slide = prs.slides.add_slide(blank_layout)
        _render_slide(pptx_slide, layout, i, manifest, resolved_template)
        _render_footer(pptx_slide, resolved_template, i + 1, len(layouts))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))

    return RenderReport(
        pptx_path=output_path,
        manifest=manifest,
        warnings=warnings,
        actual_layout=layouts,
    )


def _render_slide(
    pptx_slide: Any,
    layout: LayoutResult,
    slide_index: int,
    manifest: list[dict[str, Any]],
    template: Template,
) -> None:
    for item_id, rect in layout.rects.items():
        primitive = layout.items.get(item_id)
        if primitive is None:
            continue
        if isinstance(primitive, Shape) and primitive.kind == "connector":
            continue  # drawn below from layout.connectors, not from its own rect
        _render_primitive(
            pptx_slide, primitive, rect, slide_index, item_id, manifest, template
        )

    for points in layout.connectors.values():
        _render_connector(pptx_slide, points)


def _render_primitive(
    pptx_slide: Any,
    primitive: PrimitiveBase,
    rect: Rect,
    slide_index: int,
    item_id: str,
    manifest: list[dict[str, Any]],
    template: Template,
) -> None:
    if isinstance(primitive, Header):
        _render_header(pptx_slide, primitive, rect, template)
    elif isinstance(primitive, Text):
        _render_text(pptx_slide, primitive, rect, template)
    elif isinstance(primitive, Shape):
        _render_shape(pptx_slide, primitive, rect, template)
    elif isinstance(primitive, Image):
        _render_image(
            pptx_slide, primitive, rect, slide_index, item_id, manifest, template
        )
    elif isinstance(primitive, Stat):
        _render_stat(pptx_slide, primitive, rect, template)
    elif isinstance(primitive, Table):
        _render_table(pptx_slide, primitive, rect, template)
    elif isinstance(primitive, Sequence):
        _render_sequence(pptx_slide, primitive, rect, template)
    elif isinstance(primitive, Chart):
        _render_chart(pptx_slide, primitive, rect)
    # Grid has no visual of its own — only its (already-flattened) children render.


def _render_header(
    pptx_slide: Any, header: Header, rect: Rect, template: Template
) -> None:
    box = pptx_slide.shapes.add_textbox(
        Emu(rect.x), Emu(rect.y), Emu(rect.w), Emu(rect.h)
    )
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    first = True
    if header.eyebrow:
        p = tf.paragraphs[0]
        first = False
        p.text = header.eyebrow
        _set_font(p.font, template, 12)
        p.alignment = _ALIGN_TO_PP[header.align]

    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.text = header.title
    _set_font(p.font, template, HEADER_FONT_SIZE_PT, bold=True)
    p.alignment = _ALIGN_TO_PP[header.align]

    if header.subtitle:
        p = tf.add_paragraph()
        p.text = header.subtitle
        _set_font(p.font, template, 16)
        p.alignment = _ALIGN_TO_PP[header.align]


def _render_text(pptx_slide: Any, text: Text, rect: Rect, template: Template) -> None:
    box = pptx_slide.shapes.add_textbox(
        Emu(rect.x), Emu(rect.y), Emu(rect.w), Emu(rect.h)
    )
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    if text.mode == "paragraph":
        content = (
            text.content if isinstance(text.content, str) else " ".join(text.content)
        )
        tf.text = content
        _set_font(tf.paragraphs[0].font, template, BODY_FONT_SIZE_PT)
        return

    items = text.content if isinstance(text.content, list) else [text.content]
    emphasis = set(text.emphasis_indices or [])
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"\u2022 {item}"
        _set_font(p.font, template, BODY_FONT_SIZE_PT, bold=i in emphasis)


def _set_font(
    font: Any, template: Template, size_pt: float, *, bold: bool = False
) -> None:
    """One shared point that applies both size and typeface to a run/paragraph
    font — `template.font_family` is the agent's choice (via `Deck.template`,
    schema.py), never hardcoded per primitive.
    """
    font.size = Pt(size_pt)
    font.name = template.font_family
    if bold:
        font.bold = True


def _lighten(hex_color: str, factor: float) -> RGBColor:
    """Blend a hex color toward white by `factor` (0=unchanged, 1=white)."""
    r, g, b = (int(hex_color[i : i + 2], 16) for i in (0, 2, 4))
    return RGBColor(
        round(r + (255 - r) * factor),
        round(g + (255 - g) * factor),
        round(b + (255 - b) * factor),
    )


def _apply_shape_gradient(sp: Any, hex_color: str) -> None:
    """Blend a light tint of `hex_color` into the color itself, top to bottom.

    Only used when a shape explicitly opts in via `fill_style="gradient"` —
    the fill style is the agent's choice, not something render.py imposes.
    """
    base = hex_color.lstrip("#")
    sp.fill.gradient()
    sp.fill.gradient_angle = 90
    stops = sp.fill.gradient_stops
    stops[0].color.rgb = _lighten(base, 0.45)
    stops[1].color.rgb = RGBColor.from_string(base)


def _render_shape(
    pptx_slide: Any, shape: Shape, rect: Rect, template: Template
) -> None:
    # "line"/"arrow" render as a straight connector across the shape's own rect;
    # arrowhead styling is a known gap, left for a follow-up once needed.
    if shape.kind in ("line", "arrow"):
        pptx_slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Emu(rect.x),
            Emu(rect.y),
            Emu(rect.x + rect.w),
            Emu(rect.y + rect.h),
        )
        return

    mso_shape = _SHAPE_KIND_TO_MSO.get(shape.kind, MSO_SHAPE.RECTANGLE)
    sp = pptx_slide.shapes.add_shape(
        mso_shape, Emu(rect.x), Emu(rect.y), Emu(rect.w), Emu(rect.h)
    )

    if shape.fill and shape.fill_style == "gradient":
        _apply_shape_gradient(sp, shape.fill)
    elif shape.fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(shape.fill.lstrip("#"))
    else:
        sp.fill.background()

    if shape.border:
        sp.line.color.rgb = RGBColor.from_string(shape.border.lstrip("#"))
    else:
        sp.line.fill.background()

    if shape.text is not None:
        tf = sp.text_frame
        tf.word_wrap = True
        tf.text = shape.text.content
        tf.vertical_anchor = _VALIGN_TO_MSO[shape.text.valign]
        tf.paragraphs[0].font.name = template.font_family
        if shape.text.color:
            tf.paragraphs[0].font.color.rgb = RGBColor.from_string(
                shape.text.color.lstrip("#")
            )
        tf.paragraphs[0].alignment = _ALIGN_TO_PP[shape.text.align]


def _render_footer(
    pptx_slide: Any, template: Template, slide_number: int, total_slides: int
) -> None:
    """A consistent footer on every slide — a page number, small and out of
    the way. Reserved footer space existed since Step 2 but nothing was ever
    drawn into it, leaving every slide looking unfinished at the bottom.
    """
    x = template.margin_left
    y = template.page_height - template.margin_bottom - template.footer_height
    w = template.page_width - template.margin_left - template.margin_right

    box = pptx_slide.shapes.add_textbox(
        Emu(x), Emu(y), Emu(w), Emu(template.footer_height)
    )
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.text = f"{slide_number} / {total_slides}"
    _set_font(tf.paragraphs[0].font, template, 10)
    tf.paragraphs[0].font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    tf.paragraphs[0].alignment = PP_ALIGN.RIGHT


def _render_connector(pptx_slide: Any, points: ConnectorPoints) -> None:
    x1, y1 = points.start
    x2, y2 = points.end
    pptx_slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Emu(x1), Emu(y1), Emu(x2), Emu(y2)
    )


def _render_image(
    pptx_slide: Any,
    image: Image,
    rect: Rect,
    slide_index: int,
    item_id: str,
    manifest: list[dict[str, Any]],
    template: Template,
) -> None:
    if image.placeholder or not image.src:
        _render_image_placeholder(
            pptx_slide, image, rect, slide_index, item_id, manifest, template
        )
        return

    if image.fit == "cover":
        pptx_slide.shapes.add_picture(
            image.src, Emu(rect.x), Emu(rect.y), width=Emu(rect.w), height=Emu(rect.h)
        )
        return

    # contain: preserve aspect ratio, center within rect.
    with PILImage.open(image.src) as im:
        img_w, img_h = im.size
    img_ratio = img_w / img_h
    box_ratio = rect.w / rect.h
    if img_ratio > box_ratio:
        draw_w, draw_h = rect.w, round(rect.w / img_ratio)
    else:
        draw_h, draw_w = rect.h, round(rect.h * img_ratio)
    x = rect.x + (rect.w - draw_w) // 2
    y = rect.y + (rect.h - draw_h) // 2
    pptx_slide.shapes.add_picture(
        image.src, Emu(x), Emu(y), width=Emu(draw_w), height=Emu(draw_h)
    )


def _render_image_placeholder(
    pptx_slide: Any,
    image: Image,
    rect: Rect,
    slide_index: int,
    item_id: str,
    manifest: list[dict[str, Any]],
    template: Template,
) -> None:
    """A first-class placeholder (COMPONO_PLAN.md section 5): dashed border + caption,
    plus a manifest entry a later image-fill pass can use without re-laying-out.
    """
    sp = pptx_slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(rect.x), Emu(rect.y), Emu(rect.w), Emu(rect.h)
    )
    sp.fill.background()
    sp.line.color.rgb = RGBColor(0x99, 0x99, 0x99)
    sp.line.dash_style = MSO_LINE_DASH_STYLE.DASH

    if image.caption:
        tf = sp.text_frame
        tf.word_wrap = True
        tf.text = image.caption
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        _set_font(tf.paragraphs[0].font, template, CAPTION_FONT_SIZE_PT)
        # Auto-shapes with no fill otherwise inherit a near-invisible theme text
        # color on some renderers (observed as near-white on white) — match the
        # dashed border's gray explicitly so the caption is always legible.
        tf.paragraphs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    manifest.append(
        {
            "slide": slide_index,
            "primitive": item_id,
            "rect": {"x": rect.x, "y": rect.y, "w": rect.w, "h": rect.h},
            "caption": image.caption,
        }
    )


def _render_stat(pptx_slide: Any, stat: Stat, rect: Rect, template: Template) -> None:
    box = pptx_slide.shapes.add_textbox(
        Emu(rect.x), Emu(rect.y), Emu(rect.w), Emu(rect.h)
    )
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    p_value = tf.paragraphs[0]
    p_value.text = stat.value
    _set_font(p_value.font, template, STAT_VALUE_FONT_SIZE_PT, bold=True)
    p_value.alignment = PP_ALIGN.CENTER

    p_label = tf.add_paragraph()
    p_label.text = stat.label
    _set_font(p_label.font, template, STAT_LABEL_FONT_SIZE_PT)
    p_label.alignment = PP_ALIGN.CENTER

    if stat.trend:
        p_trend = tf.add_paragraph()
        p_trend.text = stat.trend
        _set_font(p_trend.font, template, STAT_LABEL_FONT_SIZE_PT)
        p_trend.alignment = PP_ALIGN.CENTER


def _render_table(
    pptx_slide: Any, table: Table, rect: Rect, template: Template
) -> None:
    n_rows = len(table.rows) + 1
    n_cols = len(table.headers)
    graphic_frame = pptx_slide.shapes.add_table(
        n_rows, n_cols, Emu(rect.x), Emu(rect.y), Emu(rect.w), Emu(rect.h)
    )
    tbl = graphic_frame.table

    # add_table's graphicFrame reports height=rect.h, but individual row
    # heights default to a content-based minimum, leaving visible dead space
    # below the table when it has few rows. Stretch rows to fill rect.h.
    row_h = rect.h // n_rows
    for row in tbl.rows:
        row.height = Emu(row_h)

    for c, header_text in enumerate(table.headers):
        cell = tbl.cell(0, c)
        cell.text = header_text
        _set_font(
            cell.text_frame.paragraphs[0].font, template, TABLE_FONT_SIZE_PT, bold=True
        )

    for r, row in enumerate(table.rows, start=1):
        for c, value in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = value
            is_emphasis = (
                table.emphasis_row is not None and r - 1 == table.emphasis_row
            ) or (table.emphasis_col is not None and c == table.emphasis_col)
            _set_font(
                cell.text_frame.paragraphs[0].font,
                template,
                TABLE_FONT_SIZE_PT,
                bold=is_emphasis,
            )


def _render_sequence(
    pptx_slide: Any, sequence: Sequence, rect: Rect, template: Template
) -> None:
    """Compiles internally to shape + text + connectors (COMPONO_PLAN.md section 5), not a bespoke render path."""
    steps = sequence.steps
    n = len(steps)
    if n == 0:
        return

    gutter = max(1, min(rect.w, rect.h) // 20)

    if sequence.orientation == "horizontal":
        step_w = (rect.w - gutter * (n - 1)) // n
        positions = [
            Rect(rect.x + i * (step_w + gutter), rect.y, step_w, rect.h)
            for i in range(n)
        ]
    else:
        step_h = (rect.h - gutter * (n - 1)) // n
        positions = [
            Rect(rect.x, rect.y + i * (step_h + gutter), rect.w, step_h)
            for i in range(n)
        ]

    for step, step_rect in zip(steps, positions, strict=True):
        sp = pptx_slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Emu(step_rect.x),
            Emu(step_rect.y),
            Emu(step_rect.w),
            Emu(step_rect.h),
        )
        tf = sp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.text = step.label
        _set_font(tf.paragraphs[0].font, template, BODY_FONT_SIZE_PT, bold=True)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        if step.description:
            p = tf.add_paragraph()
            p.text = step.description
            _set_font(p.font, template, CAPTION_FONT_SIZE_PT)
            p.alignment = PP_ALIGN.CENTER

    # Connect box *edges* through the gutter only — never box centers, which
    # would draw the connector across the step's own label text.
    for rect_a, rect_b in pairwise(positions):
        if sequence.orientation == "horizontal":
            x1, y1 = rect_a.x + rect_a.w, rect_a.y + rect_a.h // 2
            x2, y2 = rect_b.x, rect_b.y + rect_b.h // 2
        else:
            x1, y1 = rect_a.x + rect_a.w // 2, rect_a.y + rect_a.h
            x2, y2 = rect_b.x + rect_b.w // 2, rect_b.y
        pptx_slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Emu(x1), Emu(y1), Emu(x2), Emu(y2)
        )


def _render_chart(pptx_slide: Any, chart: Chart, rect: Rect) -> None:
    chart_data = CategoryChartData()
    chart_data.categories = chart.categories
    for series in chart.series:
        chart_data.add_series(series.name, series.values)

    pptx_slide.shapes.add_chart(
        _CHART_TYPE_TO_XL[chart.chart_type],
        Emu(rect.x),
        Emu(rect.y),
        Emu(rect.w),
        Emu(rect.h),
        chart_data,
    )
