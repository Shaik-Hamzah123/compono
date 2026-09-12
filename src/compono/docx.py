"""Orchestrates validate -> write docx.

Exposes the two public verbs for compono's second output format:
render_docx(spec, output_path) and validate_docx(spec). Mirrors render.py's
pptx pipeline in shape (parse -> per-primitive render -> report), but with
no resolver step — Word flows content top-to-bottom on its own, so there's
no EMU layout math to do for this format.

Full v1 docx catalog: heading, paragraph, bullet_list, numbered_list,
table, image, chart, page_break (docx_schema.py).

This is the only module allowed to import `docx`/`matplotlib`, mirroring
render.py's "only module allowed to import pptx" rule.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import docx as _python_docx
import matplotlib

matplotlib.use("Agg")  # headless — no GUI backend, never opens a window
import matplotlib.pyplot as plt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from pydantic import ValidationError
from pydantic_core import ErrorDetails

from compono.docx_schema import (
    BulletList,
    DocChart,
    DocImage,
    DocTable,
    DocxDoc,
    DocxPrimitiveSpec,
    Heading,
    NumberedList,
    PageBreak,
    Paragraph,
    Run,
    Section,
)
from compono.render import ValidationReport

_CHART_COLOR = "#2A6FDB"


class DocxValidationError(Exception):
    """The one exception type for docx, carrying the same structured error shape."""

    def __init__(self, errors: list[dict[str, Any]]) -> None:
        super().__init__(f"{len(errors)} validation error(s)")
        self.errors = errors


@dataclass
class DocxRenderReport:
    """render_docx returns this, not just a file — same feedback-channel intent as RenderReport."""

    docx_path: Path
    manifest: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _pydantic_error_to_dict(err: ErrorDetails) -> dict[str, Any]:
    loc = ".".join(str(part) for part in err["loc"]) or "<docx>"
    return {
        "section": None,
        "primitive": loc,
        "field": err["loc"][-1] if err["loc"] else None,
        "error": err["type"],
        "detail": err["msg"],
        "fix": "Check the field against the schema description and correct the value/type.",
    }


def _parse_doc(spec: dict[str, Any] | DocxDoc) -> DocxDoc:
    if isinstance(spec, DocxDoc):
        return spec
    try:
        return DocxDoc.model_validate(spec)
    except ValidationError as exc:
        raise DocxValidationError(
            [_pydantic_error_to_dict(e) for e in exc.errors()]
        ) from exc


def validate_docx(spec: dict[str, Any] | DocxDoc) -> ValidationReport:
    """Validate a compono docx spec: schema checks only, no file write.

    Never raises: malformed input comes back as valid=False with structured
    errors, each shaped {section, primitive, field, error, detail, fix}.
    """
    try:
        _parse_doc(spec)
    except DocxValidationError as exc:
        return ValidationReport(valid=False, errors=exc.errors)
    return ValidationReport(valid=True)


def _apply_run(paragraph: Any, run_spec: Run) -> None:
    if run_spec.link:
        _add_hyperlink(paragraph, run_spec)
        return
    run = paragraph.add_run(run_spec.text)
    run.bold = run_spec.bold
    run.italic = run_spec.italic
    run.underline = run_spec.underline


def _add_hyperlink(paragraph: Any, run_spec: Run) -> None:
    # python-docx has no add_hyperlink API — build the relationship + XML
    # by hand, the documented workaround for this exact gap.
    from docx.oxml.ns import qn
    from docx.oxml.shared import OxmlElement

    part = paragraph.part
    r_id = part.relate_to(
        run_spec.link,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if run_spec.bold:
        rpr.append(OxmlElement("w:b"))
    if run_spec.italic:
        rpr.append(OxmlElement("w:i"))
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    new_run.append(rpr)

    text_el = OxmlElement("w:t")
    text_el.text = run_spec.text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _render_heading(doc: Any, primitive: Heading) -> None:
    doc.add_heading(primitive.text, level=primitive.level)


def _render_paragraph(doc: Any, primitive: Paragraph) -> None:
    p = doc.add_paragraph()
    for run_spec in primitive.runs:
        _apply_run(p, run_spec)


def _render_bullet_list(doc: Any, primitive: BulletList) -> None:
    for item_runs in primitive.items:
        p = doc.add_paragraph(style="List Bullet")
        for run_spec in item_runs:
            _apply_run(p, run_spec)


def _render_numbered_list(doc: Any, primitive: NumberedList) -> None:
    for item_runs in primitive.items:
        p = doc.add_paragraph(style="List Number")
        for run_spec in item_runs:
            _apply_run(p, run_spec)


def _render_table(doc: Any, primitive: DocTable) -> None:
    table = doc.add_table(rows=1, cols=len(primitive.headers))
    table.style = "Light Grid Accent 1"
    header_cells = table.rows[0].cells
    for cell, text in zip(header_cells, primitive.headers, strict=True):
        cell.text = text
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for row_values in primitive.rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row_values, strict=True):
            cell.text = text


def _render_image(doc: Any, primitive: DocImage) -> list[dict[str, Any]]:
    manifest_entries: list[dict[str, Any]] = []
    if primitive.placeholder:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"[image placeholder: {primitive.caption or 'untitled'}]")
        run.italic = True
        manifest_entries.append(
            {
                "primitive": "image",
                "caption": primitive.caption,
                "width_in": primitive.width_in,
            }
        )
    else:
        doc.add_picture(primitive.src, width=Inches(primitive.width_in))
        last_p = doc.paragraphs[-1]
        last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if primitive.caption and not primitive.placeholder:
        caption_p = doc.add_paragraph(primitive.caption)
        caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in caption_p.runs:
            r.italic = True
    return manifest_entries


def _render_chart(doc: Any, primitive: DocChart) -> None:
    # No native Word chart API exists in python-docx (see DocChart's
    # docstring) — rasterize via matplotlib and embed as a picture. This is
    # a documented, deliberate exception to compono's usual "real, editable
    # object" preference: there is no editable native chart primitive to be
    # faithful to in the docx format.
    fig, ax = plt.subplots(figsize=(primitive.width_in, primitive.width_in * 0.6))
    if primitive.chart_type == "pie":
        ax.pie(
            primitive.series[0].values, labels=primitive.categories, autopct="%1.0f%%"
        )
    else:
        x = range(len(primitive.categories))
        for series in primitive.series:
            if primitive.chart_type == "bar":
                ax.bar(
                    x, series.values, label=series.name, color=_CHART_COLOR, alpha=0.85
                )
            else:
                ax.plot(
                    x, series.values, label=series.name, marker="o", color=_CHART_COLOR
                )
        ax.set_xticks(list(x))
        ax.set_xticklabels(primitive.categories)
        if len(primitive.series) > 1:
            ax.legend()
    fig.tight_layout()

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    buf.seek(0)

    doc.add_picture(buf, width=Inches(primitive.width_in))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def _render_page_break(doc: Any, primitive: PageBreak) -> None:
    doc.add_page_break()


def _render_primitive(doc: Any, primitive: DocxPrimitiveSpec) -> list[dict[str, Any]]:
    if isinstance(primitive, Heading):
        _render_heading(doc, primitive)
    elif isinstance(primitive, Paragraph):
        _render_paragraph(doc, primitive)
    elif isinstance(primitive, BulletList):
        _render_bullet_list(doc, primitive)
    elif isinstance(primitive, NumberedList):
        _render_numbered_list(doc, primitive)
    elif isinstance(primitive, DocTable):
        _render_table(doc, primitive)
    elif isinstance(primitive, DocImage):
        return _render_image(doc, primitive)
    elif isinstance(primitive, DocChart):
        _render_chart(doc, primitive)
    elif isinstance(primitive, PageBreak):
        _render_page_break(doc, primitive)
    return []


def _render_section(
    doc: Any, section: Section, *, is_first: bool
) -> list[dict[str, Any]]:
    if not is_first:
        doc.add_section()
    docx_section = doc.sections[-1]
    if section.header_text is not None:
        # A new section's header is linked to the previous one by default —
        # unlink it first, or setting text here would silently overwrite
        # every earlier section's header too.
        docx_section.header.is_linked_to_previous = False
        docx_section.header.paragraphs[0].text = section.header_text
    if section.footer_text is not None:
        docx_section.footer.is_linked_to_previous = False
        docx_section.footer.paragraphs[0].text = section.footer_text

    manifest: list[dict[str, Any]] = []
    for primitive in section.body:
        manifest.extend(_render_primitive(doc, primitive))
    return manifest


def render_docx(
    spec: dict[str, Any] | DocxDoc, output_path: str | Path
) -> DocxRenderReport:
    """Render a compono docx spec to a real, editable .docx file at output_path.

    On success returns {docx_path, manifest, warnings}. On any validation
    error, raises DocxValidationError (structured {section, primitive,
    field, error, detail, fix} errors) — nothing is written in that case.
    """
    doc_spec = _parse_doc(spec)

    doc = _python_docx.Document()
    doc.core_properties.title = doc_spec.title

    manifest: list[dict[str, Any]] = []
    for i, section in enumerate(doc_spec.sections):
        manifest.extend(_render_section(doc, section, is_first=(i == 0)))

    output = Path(output_path)
    doc.save(str(output))

    return DocxRenderReport(docx_path=output, manifest=manifest, warnings=[])
